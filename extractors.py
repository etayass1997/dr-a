# -*- coding: utf-8 -*-
"""File text/image extraction shared by app.py (/upload) and ingest_cli.py.

Kept separate from app.py so the terminal ingestion script doesn't need
Flask or an ANTHROPIC_API_KEY just to read files into materials_store.
"""
import openpyxl
import PyPDF2
from docx import Document

import materials_store

IMAGE_EXT = {'png', 'jpg', 'jpeg'}
TEXT_EXT = {'pdf', 'docx', 'xlsx', 'txt'}


def extract_pdf_text(stream):
    reader = PyPDF2.PdfReader(stream)
    parts = []
    for page in reader.pages:
        text = page.extract_text() or ''
        if text.strip():
            parts.append(text)
    return '\n'.join(parts)


def extract_docx_text(stream):
    doc = Document(stream)
    return '\n'.join(p.text for p in doc.paragraphs if p.text.strip())


def extract_xlsx_text(stream):
    wb = openpyxl.load_workbook(stream, data_only=True)
    lines = []
    for sheet in wb.worksheets:
        lines.append(f'== גיליון: {sheet.title} ==')
        for row in sheet.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None]
            if cells:
                lines.append(' | '.join(cells))
    return '\n'.join(lines)


def extract_txt_text(stream):
    raw = stream.read()
    for encoding in ('utf-8', 'cp1255', 'latin-1'):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode('utf-8', errors='ignore')


EXTRACTORS = {
    'pdf': extract_pdf_text,
    'docx': extract_docx_text,
    'xlsx': extract_xlsx_text,
    'txt': extract_txt_text,
}


def ingest_file(filename, stream):
    """Extracts+stores one file via materials_store. Returns the stored record."""
    ext = filename.rsplit('.', 1)[-1].lower() if '.' in filename else ''

    if ext in IMAGE_EXT:
        media_type = 'image/jpeg' if ext in ('jpg', 'jpeg') else 'image/png'
        try:
            return materials_store.add_material(
                filename=filename, mtype='image', media_type=media_type, image_bytes=stream.read(),
            )
        except Exception as e:
            return materials_store.add_material(filename=filename, mtype='error', error=str(e))

    extractor = EXTRACTORS.get(ext)
    if not extractor:
        return materials_store.add_material(filename=filename, mtype='error', error='סוג קובץ לא נתמך')

    try:
        text = extractor(stream)
        return materials_store.add_material(filename=filename, mtype='text', content=text)
    except Exception as e:
        return materials_store.add_material(filename=filename, mtype='error', error=str(e))
